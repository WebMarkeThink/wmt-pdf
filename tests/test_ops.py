"""Fase 2: operazioni, ottimizzazione, filigrane, annotazioni, conversioni, esportazione."""

import json
from email.message import EmailMessage

import pikepdf
import pypdfium2 as pdfium
import pytest
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from wmt_pdf.cli import main


def run(capsys, *argv):
    code = main([str(a) for a in argv])
    out = capsys.readouterr()
    return code, out.out, out.err


def page_count(path) -> int:
    with pikepdf.open(path) as pdf:
        return len(pdf.pages)


def text_of(path, page=None) -> str:
    doc = pdfium.PdfDocument(str(path))
    pages = [page] if page else range(1, len(doc) + 1)
    out = []
    for n in pages:
        tp = doc[n - 1].get_textpage()
        out.append(tp.get_text_range())
    doc.close()
    return "\n".join(out)


def numbered_pdf(path, pages: int):
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(1, pages + 1):
        c.setFont("Helvetica", 20)
        c.drawString(100, 700, f"Contenuto pagina {i}")
        c.showPage()
    c.save()
    return path


# --- operazioni di base -------------------------------------------------------


def test_merge_pdfs_with_bookmarks(capsys, tmp_path, text_pdf, table_pdf):
    out = tmp_path / "unito.pdf"
    code, stdout, _ = run(capsys, "merge", text_pdf, table_pdf, "-o", out)
    assert code == 0
    assert page_count(out) == 4
    with pikepdf.open(out) as pdf, pdf.open_outline() as outline:
        titles = [item.title for item in outline.root]
    assert titles == ["testo", "tabella"]
    assert "4 pagine" in stdout


def test_merge_unwraps_p7m(capsys, tmp_path, p7m_file, table_pdf):
    out = tmp_path / "unito.pdf"
    code, _, _ = run(capsys, "merge", p7m_file, table_pdf, "-o", out)
    assert code == 0
    assert page_count(out) == 4


def test_split_every(capsys, tmp_path):
    src = numbered_pdf(tmp_path / "cinque.pdf", 5)
    outdir = tmp_path / "parti"
    code, _, _ = run(capsys, "split", src, "--every", "2", "-o", outdir)
    assert code == 0
    parts = sorted(outdir.glob("*.pdf"))
    assert [p.name for p in parts] == ["cinque_1-2.pdf", "cinque_3-4.pdf", "cinque_5.pdf"]
    assert "Contenuto pagina 3" in text_of(parts[1], 1)


def test_split_ranges(capsys, tmp_path):
    src = numbered_pdf(tmp_path / "cinque.pdf", 5)
    outdir = tmp_path / "parti"
    run(capsys, "split", src, "--ranges", "1;2-4", "-o", outdir)
    assert sorted(p.name for p in outdir.glob("*.pdf")) == ["cinque_1.pdf", "cinque_2-4.pdf"]


def test_pages_reorder_and_delete(capsys, tmp_path):
    src = numbered_pdf(tmp_path / "tre.pdf", 3)
    out = tmp_path / "riordinato.pdf"
    run(capsys, "pages", src, "--order", "3,1", "-o", out)
    assert page_count(out) == 2
    assert "pagina 3" in text_of(out, 1)
    out2 = tmp_path / "senza2.pdf"
    run(capsys, "pages", src, "--delete", "2", "-o", out2)
    assert "pagina 2" not in text_of(out2)


def test_rotate(capsys, tmp_path):
    src = numbered_pdf(tmp_path / "due.pdf", 2)
    out = tmp_path / "ruotato.pdf"
    run(capsys, "rotate", src, "--angle", "90", "--pages", "2", "-o", out)
    with pikepdf.open(out) as pdf:
        assert int(pdf.pages[0].obj.get("/Rotate", 0)) == 0
        assert int(pdf.pages[1].obj.get("/Rotate", 0)) == 90


def test_meta(capsys, tmp_path, text_pdf):
    out = tmp_path / "meta.pdf"
    code, _, _ = run(capsys, "meta", text_pdf, "--title", "Nuovo titolo", "--author", "Studio", "-o", out)
    assert code == 0
    _, info, _ = run(capsys, "info", out)
    assert "Titolo: Nuovo titolo" in info
    assert "Autore: Studio" in info


def test_protect_and_unlock(capsys, tmp_path, text_pdf):
    locked = tmp_path / "protetto.pdf"
    code, _, _ = run(capsys, "protect", text_pdf, "--new-password", "segreta", "--no-copy", "-o", locked)
    assert code == 0
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(locked)
    code, _, err = run(capsys, "info", locked)
    assert code == 1 and "password" in err
    _, info, _ = run(capsys, "info", locked, "--password", "segreta")
    assert "Pagine: 3" in info
    unlocked = tmp_path / "libero.pdf"
    run(capsys, "unlock", locked, "--password", "segreta", "-o", unlocked)
    with pikepdf.open(unlocked) as pdf:
        assert not pdf.is_encrypted


def test_password_from_env(capsys, tmp_path, text_pdf, monkeypatch):
    monkeypatch.setenv("PDF_PW", "dallambiente")
    locked = tmp_path / "protetto.pdf"
    run(capsys, "protect", text_pdf, "--new-password-env", "PDF_PW", "-o", locked)
    with pikepdf.open(locked, password="dallambiente") as pdf:
        assert pdf.is_encrypted


def test_ops_refuse_overwrite(capsys, tmp_path, text_pdf):
    out = tmp_path / "esiste.pdf"
    out.write_bytes(b"x")
    code, _, err = run(capsys, "rotate", text_pdf, "--angle", "90", "-o", out)
    assert code == 1 and "--force" in err
    code, _, err = run(capsys, "rotate", text_pdf, "--angle", "90", "-o", text_pdf, "--force")
    assert code == 1 and "coincide" in err


# --- ottimizzazione ----------------------------------------------------------


@pytest.fixture
def heavy_pdf(tmp_path):
    """Una foto da 3000x2000 px su una pagina A4: troppo per l'email."""
    import random

    random.seed(3)
    img = Image.new("RGB", (3000, 2000))
    px = img.load()
    for x in range(0, 3000, 2):
        for y in range(0, 2000, 2):
            px[x, y] = (random.randrange(256), (x // 12) % 256, (y // 8) % 256)
    png = tmp_path / "foto.png"
    img.save(png)
    path = tmp_path / "pesante.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawImage(str(png), 30, 300, width=535, height=357)
    c.drawString(50, 100, "Didascalia della foto")
    c.showPage()
    c.save()
    return path


def test_optimize_shrinks_images(capsys, tmp_path, heavy_pdf):
    out = tmp_path / "leggero.pdf"
    code, stdout, _ = run(capsys, "optimize", heavy_pdf, "--profile", "email", "-o", out)
    assert code == 0
    assert out.stat().st_size < heavy_pdf.stat().st_size * 0.5
    assert "Didascalia della foto" in text_of(out)
    assert "%" in stdout
    with pikepdf.open(out) as pdf:
        img = next(iter(pdf.pages[0].get_images().values()))
        assert int(img.Width) <= 1300


def test_optimize_never_grows(capsys, tmp_path, text_pdf):
    out = tmp_path / "ott.pdf"
    run(capsys, "optimize", text_pdf, "-o", out)
    assert out.stat().st_size <= text_pdf.stat().st_size


# --- filigrane ---------------------------------------------------------------


def test_watermark_text_and_remove(capsys, tmp_path, text_pdf):
    marked = tmp_path / "bozza.pdf"
    code, _, _ = run(capsys, "watermark", text_pdf, "--text", "BOZZA RISERVATA", "-o", marked)
    assert code == 0
    assert "BOZZA RISERVATA" in text_of(marked, 2)
    clean = tmp_path / "pulito.pdf"
    code, stdout, _ = run(capsys, "unwatermark", marked, "-o", clean)
    assert code == 0
    assert "BOZZA" not in text_of(clean)
    assert "Milano" in text_of(clean)
    assert "3" in stdout  # filigrane tolte da 3 pagine


def test_watermark_image(capsys, tmp_path, text_pdf):
    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 100), "red").save(logo)
    out = tmp_path / "logo.pdf"
    code, _, _ = run(capsys, "watermark", text_pdf, "--image", logo, "--opacity", "0.2", "-o", out)
    assert code == 0
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages) == 3


def test_unwatermark_acrobat_style(capsys, tmp_path):
    """Filigrana come la scrive Acrobat: contenuto marcato /Artifact /Watermark nella pagina."""
    src = numbered_pdf(tmp_path / "uno.pdf", 1)
    with pikepdf.open(src, allow_overwriting_input=True) as pdf:
        page = pdf.pages[0]
        font = pikepdf.Dictionary(Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica)
        page.obj.Resources.Font.F9 = pdf.make_indirect(font)
        page.contents_add(
            pdf.make_stream(b"/Artifact <</Subtype /Watermark /Type /Pagination>> BDC BT /F9 40 Tf 100 400 Td (COPIA) Tj ET EMC"),
            prepend=False,
        )
        pdf.save(src)
    assert "COPIA" in text_of(src)
    out = tmp_path / "pulito.pdf"
    run(capsys, "unwatermark", src, "-o", out)
    assert "COPIA" not in text_of(out)
    assert "Contenuto pagina 1" in text_of(out)


# --- evidenziazioni e annotazioni -------------------------------------------


def test_highlight_and_list_annotations(capsys, tmp_path, text_pdf):
    out = tmp_path / "evidenziato.pdf"
    code, stdout, _ = run(capsys, "highlight", text_pdf, "dodici mesi", "--note", "verificare", "-o", out)
    assert code == 0
    assert "1 evidenziazion" in stdout
    with pikepdf.open(out) as pdf:
        annots = pdf.pages[1].obj.Annots
        assert annots[0].Subtype == pikepdf.Name.Highlight
        assert len(annots[0].QuadPoints) == 8
        assert "/AP" in annots[0]
    code, listing, _ = run(capsys, "annotations", out)
    assert code == 0
    assert "p.2" in listing and "Evidenziazione" in listing
    assert "dodici mesi" in listing
    assert "verificare" in listing


def test_highlight_no_match(capsys, tmp_path, text_pdf):
    out = tmp_path / "x.pdf"
    code, _stdout, _ = run(capsys, "highlight", text_pdf, "inesistente", "-o", out)
    assert code == 1
    assert not out.exists()


def test_remove_annotations(capsys, tmp_path, text_pdf):
    marked = tmp_path / "evidenziato.pdf"
    run(capsys, "highlight", text_pdf, "Milano", "-o", marked)
    out = tmp_path / "senza.pdf"
    code, _, _ = run(capsys, "annotations", marked, "--remove", "-o", out)
    assert code == 0
    _, listing, _ = run(capsys, "annotations", out)
    assert "Nessuna annotazione" in listing


# --- conversioni da altri formati ------------------------------------------------


@pytest.fixture
def docx_file(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Relazione tecnica", level=1)
    d.add_paragraph("Il sopralluogo è avvenuto a Milano il 3 ottobre.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Voce", "Importo"
    table.cell(1, 0).text, table.cell(1, 1).text = "Perizia", "800 €"
    img = tmp_path / "foto.png"
    Image.new("RGB", (2400, 1600), "blue").save(img)  # più larga della pagina
    d.add_picture(str(img))
    path = tmp_path / "relazione.docx"
    d.save(path)
    return path


@pytest.fixture
def xlsx_file(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Preventivo"
    ws.append(["Voce", "Quantità", "Prezzo"])
    ws.append(["Sito web", 1, 1500.5])
    ws.append(["Totale", None, "=C2"])
    path = tmp_path / "preventivo.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def pptx_file(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Proposta commerciale"
    slide.placeholders[1].text = "Primo punto\nSecondo punto"
    img = tmp_path / "grafico.png"
    Image.new("RGB", (400, 300), "green").save(img)
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Risultati"
    slide2.shapes.add_picture(str(img), Inches(1), Inches(2))
    path = tmp_path / "proposta.pptx"
    prs.save(path)
    return path


@pytest.fixture
def odt_file(tmp_path):
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P

    doc = OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text="Verbale di assemblea"))
    doc.text.addElement(P(text="Presenti: tutti i condomini."))
    path = tmp_path / "verbale.odt"
    doc.save(str(path))
    return path


@pytest.fixture
def eml_file(tmp_path, text_pdf):
    msg = EmailMessage()
    msg["From"] = "Studio <studio@example.com>"
    msg["To"] = "cliente@example.com"
    msg["Subject"] = "Invio documenti firmati"
    msg["Date"] = "Fri, 16 Oct 2026 10:00:00 +0200"
    msg.set_content("Buongiorno,\nin allegato il contratto.\n\nCordiali saluti")
    msg.add_attachment(text_pdf.read_bytes(), maintype="application", subtype="pdf", filename="contratto.pdf")
    path = tmp_path / "messaggio.eml"
    path.write_bytes(bytes(msg))
    return path


@pytest.mark.parametrize(
    "fixture,expected",
    [
        ("docx_file", ["Relazione tecnica", "Milano", "Perizia", "800"]),
        ("xlsx_file", ["Preventivo", "Quantità", "Sito web", "1.500,50"]),
        ("pptx_file", ["Proposta commerciale", "Secondo punto", "Risultati"]),
        ("odt_file", ["Verbale di assemblea", "condomini"]),
    ],
)
def test_convert_office(capsys, request, tmp_path, fixture, expected):
    src = request.getfixturevalue(fixture)
    out = tmp_path / "convertito.pdf"
    code, _, err = run(capsys, "convert", src, "-o", out)
    assert code == 0, err
    text = text_of(out)
    for word in expected:
        assert word in text, (word, text[:300])


def test_convert_pptx_one_page_per_slide(capsys, tmp_path, pptx_file):
    out = tmp_path / "slide.pdf"
    run(capsys, "convert", pptx_file, "-o", out)
    assert page_count(out) == 2
    with pikepdf.open(out) as pdf:
        box = pdf.pages[0].mediabox
        assert float(box[2]) > float(box[3])  # orizzontale


def test_convert_text_formats(capsys, tmp_path):
    md = tmp_path / "note.md"
    md.write_text("# Titolo nota\n\n| A | B |\n|---|---|\n| uno | due |\n\nTesto con **grassetto** e città.", encoding="utf-8")
    txt = tmp_path / "appunti.txt"
    txt.write_text("Riga uno\nRiga due con è", encoding="utf-8")
    html = tmp_path / "pagina.html"
    html.write_text("<h1>Pagina web</h1><p>Paragrafo <b>uno</b></p><img src='http://example.com/x.png'>", encoding="utf-8")
    csvf = tmp_path / "dati.csv"
    csvf.write_text("nome;importo\nRossi;10\n", encoding="utf-8")
    for src, word in [(md, "Titolo nota"), (txt, "Riga due con è"), (html, "Pagina web"), (csvf, "Rossi")]:
        out = tmp_path / f"{src.stem}.pdf"
        code, _, err = run(capsys, "convert", src, "-o", out)
        assert code == 0, err
        assert word in text_of(out)


def test_convert_images(capsys, tmp_path):
    jpg = tmp_path / "foto.jpg"
    Image.new("RGB", (1200, 800), "orange").save(jpg, dpi=(300, 300))
    tif = tmp_path / "scansioni.tif"
    frames = [Image.new("L", (850, 1100), v) for v in (200, 100)]
    frames[0].save(tif, save_all=True, append_images=frames[1:], dpi=(100, 100))
    out = tmp_path / "foto.pdf"
    run(capsys, "convert", jpg, "-o", out)
    with pikepdf.open(out) as pdf:
        box = pdf.pages[0].mediabox
        assert round(float(box[2])) == 288  # 1200 px a 300 dpi = 4 pollici
    out2 = tmp_path / "scansioni.pdf"
    run(capsys, "convert", tif, "-o", out2)
    assert page_count(out2) == 2


def test_convert_email_with_attachments(capsys, tmp_path, eml_file):
    out = tmp_path / "email.pdf"
    code, _stdout, _ = run(capsys, "convert", eml_file, "--attachments", "-o", out)
    assert code == 0
    text = text_of(out)
    assert "Invio documenti firmati" in text
    assert "contratto.pdf" in text  # elencato tra gli allegati
    assert "Milano" in text  # e accodato
    assert page_count(out) == 4


def test_merge_mixed_formats(capsys, tmp_path, docx_file, xlsx_file, text_pdf):
    jpg = tmp_path / "foto.jpg"
    Image.new("RGB", (600, 400), "orange").save(jpg)
    out = tmp_path / "fascicolo.pdf"
    code, _, err = run(capsys, "merge", docx_file, xlsx_file, jpg, text_pdf, "-o", out)
    assert code == 0, err
    text = text_of(out)
    assert "Relazione tecnica" in text and "Sito web" in text and "Milano" in text
    with pikepdf.open(out) as pdf, pdf.open_outline() as outline:
        assert [i.title for i in outline.root] == ["relazione", "preventivo", "foto", "testo"]


def test_convert_unsupported(capsys, tmp_path):
    f = tmp_path / "archivio.rar"
    f.write_bytes(b"Rar!")
    code, _, err = run(capsys, "convert", f, "-o", tmp_path / "x.pdf")
    assert code == 1 and "non supportato" in err


# --- esportazione ---------------------------------------------------------------


def test_export_docx(capsys, tmp_path, table_pdf):
    import docx

    out = tmp_path / "tabella.docx"
    code, _, _ = run(capsys, "export", table_pdf, "--to", "docx", "-o", out)
    assert code == 0
    d = docx.Document(str(out))
    assert any("Preventivo" in p.text for p in d.paragraphs)
    assert d.tables[0].cell(1, 0).text == "Sito web"


def test_export_xlsx(capsys, tmp_path, table_pdf):
    import openpyxl

    out = tmp_path / "tabelle.xlsx"
    run(capsys, "export", table_pdf, "--to", "xlsx", "-o", out)
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == "Sito web"
    assert ws["C2"].value == 1500  # numero, non testo


def test_export_html_md_txt(capsys, tmp_path, table_pdf):
    for fmt, needle in [("html", "<table"), ("md", "| Voce |"), ("txt", "Preventivo")]:
        out = tmp_path / f"out.{fmt}"
        code, _, _ = run(capsys, "export", table_pdf, "--to", fmt, "-o", out)
        assert code == 0
        assert needle in out.read_text(encoding="utf-8")


def test_export_images(capsys, tmp_path, text_pdf):
    outdir = tmp_path / "img"
    code, _, _ = run(capsys, "export", text_pdf, "--to", "png", "--dpi", "40", "-o", outdir)
    assert code == 0
    assert sorted(p.name for p in outdir.glob("*.png")) == ["testo_p1.png", "testo_p2.png", "testo_p3.png"]
    tif = tmp_path / "tutto.tiff"
    run(capsys, "export", text_pdf, "--to", "tiff", "--dpi", "40", "-o", tif)
    with Image.open(tif) as im:
        assert im.n_frames == 3


def test_italian_numbers():
    from wmt_pdf.export import parse_number

    assert parse_number("1.500,50") == 1500.5
    assert parse_number("1500") == 1500
    assert parse_number("12,5%") == "12,5%"
    assert parse_number("00123") == "00123"  # codici, non numeri
    assert parse_number("Sito") == "Sito"


# --- PDF ricercabile -------------------------------------------------------------


@pytest.mark.ocr
def test_ocr_makes_searchable_pdf(capsys, tmp_path, scanned_pdf):
    pytest.importorskip("rapidocr")
    out = tmp_path / "ricercabile.pdf"
    code, _stdout, _ = run(capsys, "ocr", scanned_pdf, "-o", out)
    assert code == 0
    assert "4521" in text_of(out)
    _, found, _ = run(capsys, "search", out, "fattura")
    assert "1 risultato" in found
    # l'immagine resta quella di prima: il testo è invisibile
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages[0].get_images()) >= 1


def test_json_outputs(capsys, tmp_path, text_pdf):
    out = tmp_path / "h.pdf"
    run(capsys, "highlight", text_pdf, "Milano", "-o", out)
    _, listing, _ = run(capsys, "annotations", out, "--json")
    data = json.loads(listing)
    assert data[0]["type"] == "Highlight" and data[0]["page"] == 1


# --- sicurezza delle conversioni ------------------------------------------------------


def _png(path, color="red"):
    Image.new("RGB", (60, 40), color).save(path)
    return path


def test_html_cannot_read_files_outside_its_folder(capsys, tmp_path):
    secret = _png(tmp_path / "segreto.png")
    docs = tmp_path / "docs"
    docs.mkdir()
    _png(docs / "vicina.png", "blue")
    html = docs / "pagina.html"
    html.write_text(
        f"<p>Prova</p><img src='vicina.png'><img src='../segreto.png'><img src='{secret.as_posix()}'>"
        f"<img src='file:///{secret.as_posix()}'>",
        encoding="utf-8",
    )
    out = tmp_path / "pagina.pdf"
    code, _, _ = run(capsys, "convert", html, "-o", out)
    assert code == 0
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages[0].get_images()) == 1  # solo l'immagine nella sua cartella


def test_email_html_cannot_read_local_files(capsys, tmp_path):
    secret = _png(tmp_path / "segreto.png")
    msg = EmailMessage()
    msg["Subject"] = "Prova"
    msg.set_content("testo")
    msg.add_alternative(f"<p>Ciao</p><img src='{secret.as_posix()}'><img src='segreto.png'>", subtype="html")
    eml = tmp_path / "m.eml"
    eml.write_bytes(bytes(msg))
    out = tmp_path / "m.pdf"
    assert run(capsys, "convert", eml, "-o", out)[0] == 0
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages[0].get_images()) == 0


def test_html_does_not_fetch_remote(capsys, tmp_path, monkeypatch):
    import urllib.request

    calls = []
    monkeypatch.setattr(urllib.request, "urlopen", lambda *a, **k: calls.append(a) or (_ for _ in ()).throw(OSError()))
    html = tmp_path / "p.html"
    html.write_text("<p>Testo</p><img src='http://127.0.0.1:1/x.png'><img src='https://example.com/y.png'>",
                    encoding="utf-8")
    assert run(capsys, "convert", html, "-o", tmp_path / "p.pdf")[0] == 0
    assert calls == []


# --- regressioni dalla revisione della fase 2 ------------------------------------------


def test_unc_paths_are_never_touched(capsys, tmp_path, monkeypatch):
    import os

    touched = []
    real = os.path.realpath
    monkeypatch.setattr(os.path, "realpath", lambda p, *a, **k: touched.append(str(p)) or real(p, *a, **k))
    html = tmp_path / "unc.html"
    html.write_text(r"<p>x</p><img src='//127.0.0.9/share/a.png'><img src='\\127.0.0.9\share\b.png'>",
                    encoding="utf-8")
    code, _, err = run(capsys, "convert", html, "-o", tmp_path / "unc.pdf")
    assert code == 0, err
    assert not any("127.0.0.9" in t for t in touched)


def test_write_commands_keep_encryption(capsys, tmp_path, text_pdf):
    locked = tmp_path / "protetto.pdf"
    run(capsys, "protect", text_pdf, "--new-password", "pw", "-o", locked)
    for cmd, extra in [("rotate", ["--angle", "90"]), ("pages", ["--keep", "1-2"]), ("watermark", ["--text", "X"]),
                       ("optimize", [])]:
        out = tmp_path / f"{cmd}.pdf"
        code, _, err = run(capsys, cmd, locked, *extra, "--password", "pw", "-o", out)
        assert code == 0, err
        with pikepdf.open(out, password="pw") as pdf:
            assert pdf.is_encrypted, cmd


def test_merge_with_protected_input(capsys, tmp_path, text_pdf, table_pdf):
    locked = tmp_path / "protetto.pdf"
    run(capsys, "protect", text_pdf, "--new-password", "pw", "-o", locked)
    out = tmp_path / "unito.pdf"
    code, _, err = run(capsys, "merge", locked, table_pdf, "--password", "pw", "-o", out)
    assert code == 0, err
    with pikepdf.open(out, password="pw") as pdf:
        assert pdf.is_encrypted and len(pdf.pages) == 4


def test_protect_restrictions_are_real(capsys, tmp_path, text_pdf):
    locked = tmp_path / "limitato.pdf"
    run(capsys, "protect", text_pdf, "--new-password", "pw", "--no-copy", "--no-print", "-o", locked)
    with pikepdf.open(locked, password="pw") as pdf:
        assert not pdf.owner_password_matched
        assert not pdf.allow.extract and not pdf.allow.print_highres


def test_huge_sheet_is_refused(capsys, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "inizio"
    wb.active["XFD1048576"] = "fine"
    src = tmp_path / "enorme.xlsx"
    wb.save(src)
    code, _, err = run(capsys, "convert", src, "-o", tmp_path / "x.pdf")
    assert code == 1 and "troppo grande" in err


def test_export_images_needs_force(capsys, tmp_path, text_pdf):
    outdir = tmp_path / "img"
    run(capsys, "export", text_pdf, "--to", "png", "--dpi", "30", "-o", outdir)
    code, _, err = run(capsys, "export", text_pdf, "--to", "png", "--dpi", "30", "-o", outdir)
    assert code == 1 and "--force" in err
    assert run(capsys, "export", text_pdf, "--to", "png", "--dpi", "30", "-o", outdir, "--force")[0] == 0


def _page_with(tmp_path, content: bytes):
    src = numbered_pdf(tmp_path / "base.pdf", 1)
    with pikepdf.open(src, allow_overwriting_input=True) as pdf:
        page = pdf.pages[0]
        font = pikepdf.Dictionary(Type=pikepdf.Name.Font, Subtype=pikepdf.Name.Type1, BaseFont=pikepdf.Name.Helvetica)
        page.obj.Resources.Font.F9 = pdf.make_indirect(font)
        page.contents_add(pdf.make_stream(content), prepend=False)
        pdf.save(src)
    return src


def test_unwatermark_keeps_inline_images(capsys, tmp_path):
    src = _page_with(tmp_path, b"/Artifact <</Subtype /Watermark>> BDC BT /F9 30 Tf 100 400 Td (COPIA) Tj ET EMC "
                               b"q 10 0 0 10 50 50 cm BI /W 1 /H 1 /BPC 8 /CS /G ID \x80 EI Q")
    out = tmp_path / "pulito.pdf"
    code, _, err = run(capsys, "unwatermark", src, "-o", out)
    assert code == 0, err
    assert "COPIA" not in text_of(out)
    with pikepdf.open(out) as pdf:
        assert b"BI" in pdf.pages[0].obj.Contents.read_bytes()


def test_unwatermark_unbalanced_leaves_page(capsys, tmp_path):
    src = _page_with(tmp_path, b"/Artifact <</Subtype /Watermark>> BDC BT /F9 30 Tf 100 400 Td (COPIA) Tj ET")
    out = tmp_path / "pulito.pdf"
    run(capsys, "unwatermark", src, "-o", out)
    assert "Contenuto pagina 1" in text_of(out)


def test_split_duplicate_ranges(capsys, tmp_path):
    src = numbered_pdf(tmp_path / "tre.pdf", 3)
    code, _, err = run(capsys, "split", src, "--ranges", "1-2;1-2", "-o", tmp_path / "parti")
    assert code == 1 and "uguali" in err


def test_unexpected_errors_are_short(capsys, tmp_path):
    bad = tmp_path / "rotto.pdf"
    bad.write_bytes(b"%PDF-1.7\n" + b"\x00" * 50)
    code, _, err = run(capsys, "annotations", bad)
    assert code == 1
    assert "Traceback" not in err and err.count("\n") <= 2


def test_xml_safe_strips_control_chars():
    from wmt_pdf.export import xml_safe

    assert xml_safe("a\x02b\x0cc\nd\te") == "abc\nd\te"


# --- esportazione strutturata ------------------------------------------------------------


def test_export_json_structure(capsys, tmp_path, table_pdf):
    out = tmp_path / "doc.json"
    code, _, err = run(capsys, "export", table_pdf, "--to", "json", "-o", out)
    assert code == 0, err
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["page_count"] == 1
    page = data["pages"][0]
    assert page["number"] == 1 and page["width"] > 500
    assert "Preventivo" in page["text"]
    assert page["tables"][0][1] == ["Sito web", "1", "1500"]
    line = next(seg for seg in page["lines"] if "Preventivo" in seg["text"])
    x0, y0, x1, y1 = line["bbox"]
    assert x0 < x1 and y0 < y1 and y1 > 780  # coordinate PDF, origine in basso a sinistra
    assert data["fields"] == [] and data["annotations"] == []


def test_export_xml_structure(capsys, tmp_path, text_pdf):
    import xml.etree.ElementTree as ET

    out = tmp_path / "doc.xml"
    code, _, err = run(capsys, "export", text_pdf, "--to", "xml", "-o", out)
    assert code == 0, err
    root = ET.parse(out).getroot()
    assert root.tag == "document" and root.get("page_count") == "3"
    pages = root.findall("page")
    assert len(pages) == 3
    assert "Milano" in pages[0].find("text").text
    assert pages[0].find("lines/line").get("bbox")
    assert root.find("metadata/item[@name='Title']").text == "Contratto di prova"
