"""Conversione in PDF da altri formati, con una libreria dedicata per ciascuno.

Si conservano contenuto e struttura (titoli, paragrafi, tabelle, immagini), non l'impaginazione
esatta dell'originale: per quella servirebbe il programma che ha creato il file.
"""

from __future__ import annotations

import base64
import csv
import email
import email.policy
import html as htmllib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from wmt_pdf.i18n import t
from wmt_pdf.source import CliError, load_source

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SUPPORTED = {
    ".pdf": "PDF", ".p7m": "PDF firmato", ".docx": "Word", ".xlsx": "Excel", ".xlsm": "Excel",
    ".pptx": "PowerPoint", ".odt": "OpenDocument", ".md": "Markdown", ".markdown": "Markdown",
    ".txt": "testo", ".html": "HTML", ".htm": "HTML", ".csv": "CSV", ".eml": "email",
} | {ext: "immagine" for ext in IMAGE_EXT}

PAGE_W_PT, PAGE_H_PT = 595.28, 841.89  # A4
MARGIN_PT = 56.7  # 2 cm
DEFAULT_IMAGE_DPI = 150


@dataclass
class Converted:
    data: bytes
    notes: list[str] = field(default_factory=list)


def kind_of(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        formats = ", ".join(sorted({s.lstrip(".") for s in SUPPORTED}))
        raise CliError(t("formato non supportato: {name} (supportati: {formats})", name=path.name, formats=formats))
    return suffix


ZIP_FORMATS = {".docx", ".xlsx", ".xlsm", ".pptx", ".odt"}
MAX_UNZIPPED_BYTES = 500 * 1024 * 1024  # un documento Office vero sta molto sotto; oltre è un archivio-bomba
MAX_SHEET_CELLS = 2_000_000


def to_pdf(path: Path, *, attachments: bool = False, local_files: bool = True) -> Converted:
    """Restituisce il PDF corrispondente al file; un PDF o un p7m con PDF passano così come sono.

    local_files=False: l'HTML non può leggere nessun file locale (per allegati e contenuti non fidati).
    """
    suffix = kind_of(path)
    if suffix in (".pdf", ".p7m"):
        return Converted(load_source(path).data)
    _require_convert_extra(suffix)
    if suffix in IMAGE_EXT:
        return Converted(images_to_pdf(path))
    raw = path.read_bytes()
    if suffix in ZIP_FORMATS:
        _check_zip(raw, path.name)
    base = path.parent if local_files else None
    if suffix == ".docx":
        return Converted(html_to_pdf(docx_to_html(raw), base))
    if suffix in (".xlsx", ".xlsm"):
        return Converted(html_to_pdf(xlsx_to_html(raw), base, landscape=True))
    if suffix == ".pptx":
        return Converted(pptx_to_pdf(raw))
    if suffix == ".odt":
        return Converted(html_to_pdf(odt_to_html(raw), base))
    if suffix in (".md", ".markdown"):
        return Converted(html_to_pdf(markdown_to_html(_decode_text(raw)), base))
    if suffix == ".txt":
        return Converted(html_to_pdf(text_to_html(_decode_text(raw)), base))
    if suffix in (".html", ".htm"):
        return Converted(html_to_pdf(_decode_text(raw), base))
    if suffix == ".csv":
        return Converted(html_to_pdf(csv_to_html(_decode_text(raw)), base, landscape=True))
    if suffix == ".eml":
        return eml_to_pdf(raw, attachments=attachments)
    raise CliError(t("formato non supportato: {name}", name=path.name))  # pragma: no cover


def _check_zip(raw: bytes, name: str) -> None:
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise CliError(t("{name} non è un documento Office valido (archivio rovinato)", name=name)) from None
    if total > MAX_UNZIPPED_BYTES:
        raise CliError(t("{name}: contenuto decompresso di {v} MB, oltre il limite di sicurezza", name=name, v=total // 1024 // 1024))


def _require_convert_extra(suffix: str) -> None:
    import importlib.util

    module = {".docx": "mammoth", ".xlsx": "openpyxl", ".xlsm": "openpyxl", ".pptx": "pptx", ".odt": "odf",
              ".md": "markdown", ".markdown": "markdown"}.get(suffix, "xhtml2pdf")
    if suffix in IMAGE_EXT:
        return
    if importlib.util.find_spec(module) is None or importlib.util.find_spec("xhtml2pdf") is None:
        raise CliError(t("conversione non installata: `uv sync --extra convert` (oppure pip install 'wmt-pdf[convert]')"))


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


# --- immagini ------------------------------------------------------------------------


def images_to_pdf(path: Path) -> bytes:
    """Una pagina per immagine (o per fotogramma di un TIFF), grande quanto l'immagine alla sua risoluzione."""
    from PIL import Image, ImageSequence
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    with Image.open(path) as img:
        multi = getattr(img, "n_frames", 1) > 1
        single_jpeg = path.suffix.lower() in (".jpg", ".jpeg") and not multi
        frames = ImageSequence.Iterator(img) if multi else [img]
        for frame in frames:
            dpi = frame.info.get("dpi", (DEFAULT_IMAGE_DPI, DEFAULT_IMAGE_DPI))
            dx, dy = (float(dpi[0]) or DEFAULT_IMAGE_DPI), (float(dpi[1]) or DEFAULT_IMAGE_DPI)
            w, h = frame.size[0] * 72 / dx, frame.size[1] * 72 / dy
            single = frame.copy()
            if single.mode not in ("RGB", "L", "1", "CMYK"):
                single = single.convert("RGB")
            # Un JPEG singolo entra così com'è, senza ricompressione.
            reader = ImageReader(str(path)) if single_jpeg else ImageReader(single)
            c.setPageSize((w, h))
            c.drawImage(reader, 0, 0, w, h)
            c.showPage()
    c.save()
    return buf.getvalue()


# --- HTML → PDF ----------------------------------------------------------------------

BASE_CSS = """
@page {{ size: {size}; margin: 2cm; }}
body {{ font-family: {font}; font-size: 10pt; line-height: 1.35; }}
h1 {{ font-size: 18pt; }} h2 {{ font-size: 14pt; }} h3 {{ font-size: 12pt; }}
table {{ width: 100%; margin: 6pt 0; }}
td, th {{ border: 0.5pt solid #888; padding: 3pt; vertical-align: top; }}
th {{ background-color: #eee; }}
pre {{ font-family: Courier; font-size: 9pt; }}
.meta td {{ border: none; padding: 1pt 6pt 1pt 0; }}
"""

_FONT_CANDIDATES = [
    "C:/Windows/Fonts/arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
]


def _font_files() -> tuple[str, str] | None:
    """Un font di sistema con tutti i caratteri (accenti, €, simboli) e il suo grassetto."""
    for candidate in _FONT_CANDIDATES:
        if Path(candidate).is_file():
            bold = candidate.replace("arial.ttf", "arialbd.ttf").replace("Arial.ttf", "Arial Bold.ttf")
            bold = bold.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            return candidate, bold if Path(bold).is_file() else candidate
    return None


def _body_font() -> tuple[str, str, list[str]]:
    """Famiglia CSS, regole @font-face e file da autorizzare; Helvetica se non c'è un font di sistema."""
    files = _font_files()
    if not files:
        return "Helvetica", "", []
    regular, bold = files
    face = f"@font-face {{ font-family: Corpo; src: url('{regular}'); }}\n"
    face += f"@font-face {{ font-family: Corpo; src: url('{bold}'); font-weight: bold; }}\n"
    return "Corpo", face, [regular, bold]


def html_to_pdf(html: str, base_dir: Path | None, landscape: bool = False) -> bytes:
    """base_dir: cartella da cui si possono leggere immagini locali; None per contenuti non fidati."""
    import logging

    from xhtml2pdf import pisa
    from xhtml2pdf.config.resources import ResourceAccessPolicy

    logging.getLogger("xhtml2pdf").setLevel(logging.ERROR)  # gli avvisi su CSS e risorse bloccate costano token

    font, face, font_paths = _body_font()
    css = face + BASE_CSS.format(size="a4 landscape" if landscape else "a4", font=font)
    width = (PAGE_H_PT if landscape else PAGE_W_PT) - 2 * MARGIN_PT
    height = (PAGE_W_PT if landscape else PAGE_H_PT) - 2 * MARGIN_PT
    html = fit_images(html, width, height * 0.9)
    if not re.search(r"<html\b", html, re.IGNORECASE):
        html = f"<html><head><meta charset='utf-8'></head><body>{html}</body></html>"
    style = f"<style>{css}</style>"
    if re.search(r"<head\b[^>]*>", html, re.IGNORECASE):
        html = re.sub(r"(<head\b[^>]*>)", lambda m: m.group(1) + style, html, count=1, flags=re.IGNORECASE)
    else:
        html = re.sub(r"(<html\b[^>]*>)", lambda m: m.group(1) + f"<head>{style}</head>", html, count=1, flags=re.IGNORECASE)

    def link_callback(uri: str, rel: str) -> str | None:
        # Niente download da internet: solo file locali accanto al documento e data URI.
        if uri.startswith("data:") or uri in font_paths:
            return uri
        # Mai None: xhtml2pdf ripiegherebbe sull'indirizzo originale (e su Windows su //server/...).
        return _local_resource(uri, base_dir) or BLOCKED

    # Politica "server" di xhtml2pdf: tutti i limiti attivi, niente rete, letture solo dove serve.
    font_dirs = tuple({Path(f).parent for f in font_paths})
    root = base_dir if base_dir is not None else (font_dirs[0] if font_dirs else None)
    # Un Word con molte foto supera i 10 MB di HTML: limite più largo solo per i documenti locali.
    larger = {"max_document_bytes": 200 * 1024 * 1024} if base_dir is not None else {}
    policy = ResourceAccessPolicy.server(base_dir=root, allow_remote=False, extra_roots=font_dirs, **larger,
                                         allow_remote_pdf_backgrounds=False)
    out = io.BytesIO()
    status = pisa.CreatePDF(html, dest=out, encoding="utf-8", link_callback=link_callback, resource_policy=policy)
    if status.err:
        raise CliError(t("conversione HTML non riuscita ({err} errori)", err=status.err))
    return out.getvalue()


def _local_resource(uri: str, base_dir: Path | None) -> str | None:
    """Percorso di un file nella cartella del documento, oppure None.

    Tutti i controlli sono sul testo, prima di toccare il disco: su Windows anche solo risolvere
    //server/cartella aprirebbe una connessione di rete (e invierebbe le credenziali).
    """
    import ntpath
    import os
    import posixpath

    if base_dir is None or not uri or len(uri) > 1024:
        return None
    if re.match(r"^[a-z][a-z0-9+.-]*:", uri, re.IGNORECASE):
        return None  # http:, file:, C:
    if uri.startswith(("/", "\\")) or ntpath.isabs(uri) or posixpath.isabs(uri) or ntpath.splitdrive(uri)[0]:
        return None  # assoluti e UNC
    parts = re.split(r"[\\/]+", uri)
    if any(p == ".." for p in parts):
        return None  # niente risalite
    local = os.path.normpath(os.path.join(str(base_dir), *parts))
    try:
        root = os.path.realpath(str(base_dir))
        real = os.path.realpath(local)
        if os.path.commonpath([root, real]) != root or not os.path.isfile(real):
            return None  # collegamenti che portano fuori dalla cartella
    except (OSError, ValueError):
        return None
    return real


BLOCKED = "data:text/plain;base64,"  # risorsa vuota al posto di quelle non permesse

_IMG_TAG = re.compile(r"<img\b[^>]*>", re.IGNORECASE)


def fit_images(html: str, max_w: float, max_h: float) -> str:
    """Dà a ogni immagine una dimensione che sta nella pagina: xhtml2pdf non ridimensiona da solo."""

    def fix(match: re.Match) -> str:
        tag = match.group(0)
        src = re.search(r"""src\s*=\s*["']([^"']+)["']""", tag, re.IGNORECASE)
        size = _image_size(src.group(1)) if src else None
        if not size:
            return tag
        w, h = size[0] * 0.75, size[1] * 0.75  # px a 96 dpi → punti
        scale = min(1.0, max_w / w, max_h / h)
        tag = re.sub(r"""\s(width|height)\s*=\s*["'][^"']*["']""", "", tag, flags=re.IGNORECASE)
        return tag[:-1].rstrip("/") + f' width="{w * scale:.0f}" height="{h * scale:.0f}">'

    return _IMG_TAG.sub(fix, html)


def _image_size(src: str) -> tuple[int, int] | None:
    from PIL import Image

    if not src.startswith("data:"):
        return None
    try:
        payload = base64.b64decode(src.split(",", 1)[1])
        with Image.open(io.BytesIO(payload)) as img:
            return img.size
    except Exception:  # noqa: BLE001 - un'immagine illeggibile resta com'è
        return None


def _esc(value) -> str:
    return htmllib.escape("" if value is None else str(value))


# --- singoli formati -----------------------------------------------------------------


def docx_to_html(raw: bytes) -> str:
    import mammoth

    result = mammoth.convert_to_html(io.BytesIO(raw), convert_image=mammoth.images.data_uri)
    return result.value


def text_to_html(text: str) -> str:
    blocks = re.split(r"\n\s*\n", text.strip())
    return "".join("<p>" + "<br/>".join(_esc(line) for line in b.split("\n")) + "</p>" for b in blocks)


def markdown_to_html(text: str) -> str:
    import markdown

    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])


def csv_to_html(text: str) -> str:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _table_html(rows, header=True)


def _table_html(rows, header: bool = True) -> str:
    if not rows:
        return ""
    out = ["<table>"]
    for i, row in enumerate(rows):
        tag = "th" if header and i == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{_esc(c)}</{tag}>" for c in row) + "</tr>")
    out.append("</table>")
    return "".join(out)


def xlsx_to_html(raw: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        cells = (ws.max_row or 0) * (ws.max_column or 0)
        if cells > MAX_SHEET_CELLS:
            wb.close()
            raise CliError(t("foglio «{title}» troppo grande ({max_row} righe × {max_column} colonne)", title=ws.title, max_row=ws.max_row, max_column=ws.max_column))
        rows = [[_xlsx_value(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c != "" for c in r)]
        if not rows:
            continue
        width = max(max((i + 1 for i, c in enumerate(r) if c != ""), default=0) for r in rows)
        rows = [r[:width] + [""] * (width - len(r)) for r in rows]
        parts.append(f"<h2>{_esc(ws.title)}</h2>" + _table_html(rows))
    wb.close()
    return "".join(parts) or "<p>(foglio vuoto)</p>"


def _xlsx_value(value) -> str:
    import datetime as dt

    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%d/%m/%Y %H:%M") if value.time() != dt.time() else value.strftime("%d/%m/%Y")
    if isinstance(value, dt.date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        text = f"{value:,.2f}" if value != int(value) else f"{int(value):,}"
        return text.replace(",", "X").replace(".", ",").replace("X", ".")  # formato italiano
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".") if abs(value) >= 10000 else str(value)
    return str(value)


def odt_to_html(raw: bytes) -> str:
    import zipfile

    from odf import draw, opendocument, table, teletype, text
    from odf.namespaces import TABLENS, TEXTNS

    doc = opendocument.load(io.BytesIO(raw))
    archive = zipfile.ZipFile(io.BytesIO(raw))
    parts: list[str] = []

    def walk(node) -> None:
        for child in node.childNodes:
            qname = getattr(child, "qname", None)
            if qname == (TEXTNS, "h"):
                level = int(child.getAttribute("outlinelevel") or 1)
                parts.append(f"<h{min(level, 6)}>{_esc(teletype.extractText(child))}</h{min(level, 6)}>")
            elif qname == (TEXTNS, "p"):
                frames = child.getElementsByType(draw.Image)
                for frame in frames:
                    href = frame.getAttribute("href")
                    if href and href in archive.namelist():
                        data = base64.b64encode(archive.read(href)).decode()
                        mime = "image/png" if href.lower().endswith(".png") else "image/jpeg"
                        parts.append(f'<img src="data:{mime};base64,{data}"/>')
                parts.append(f"<p>{_esc(teletype.extractText(child))}</p>")
            elif qname == (TEXTNS, "list"):
                items = child.getElementsByType(text.ListItem)
                parts.append("<ul>" + "".join(f"<li>{_esc(teletype.extractText(i))}</li>" for i in items) + "</ul>")
            elif qname == (TABLENS, "table"):
                rows = []
                for row in child.getElementsByType(table.TableRow):
                    rows.append([teletype.extractText(cell) for cell in row.getElementsByType(table.TableCell)])
                parts.append(_table_html(rows, header=False))
            elif getattr(child, "childNodes", None):
                walk(child)

    walk(doc.text)
    return "".join(parts)


def pptx_to_pdf(raw: bytes) -> bytes:
    """Una pagina per diapositiva, delle stesse proporzioni; testi, immagini e tabelle al loro posto."""
    from pptx import Presentation
    from pptx.util import Emu
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Frame, Paragraph

    font, bold = _reportlab_font()
    prs = Presentation(io.BytesIO(raw))
    w = Emu(prs.slide_width).pt
    h = Emu(prs.slide_height).pt
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(w, h))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            x, top = Emu(shape.left).pt, Emu(shape.top).pt
            sw, sh = Emu(shape.width or 0).pt, Emu(shape.height or 0).pt
            y = h - top - sh
            if shape.shape_type == 13 and hasattr(shape, "image"):  # immagine
                try:
                    c.drawImage(ImageReader(io.BytesIO(shape.image.blob)), x, y, sw, sh, mask="auto")
                except Exception:  # noqa: BLE001, S110 - formati d'immagine esotici (EMF, WMF): si salta
                    pass
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                _draw_table(c, rows, x, y, sw, sh, font)
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                is_title = shape == getattr(slide.shapes, "title", None)
                flow = []
                for para in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in para.runs) or para.text
                    if not txt.strip():
                        continue
                    sizes = [r.font.size.pt for r in para.runs if r.font.size]
                    size = sizes[0] if sizes else (32 if is_title else 18)
                    style = ParagraphStyle("p", fontName=bold if is_title else font, fontSize=size, leading=size * 1.2,
                                           leftIndent=para.level * 18)
                    bullet = "" if is_title else "• "
                    flow.append(Paragraph(bullet + _esc(txt), style))
                frame = Frame(x, y, max(sw, 20), max(sh, 20), showBoundary=0, leftPadding=4, rightPadding=4)
                frame.addFromList(flow, c)  # quel che non entra resta fuori, come nella diapositiva
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            c.setFont(font, 7)
            c.drawString(10, 6, ("Note: " + notes.replace("\n", " "))[:180])
        c.showPage()
    c.save()
    return buf.getvalue()


def _draw_table(c, rows, x, y, w, h, font) -> None:
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    if not rows:
        return
    t = Table(rows, colWidths=[w / max(len(r) for r in rows)] * max(len(r) for r in rows))
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("FONTNAME", (0, 0), (-1, -1), font)]))
    _, th = t.wrapOn(c, w, h)
    t.drawOn(c, x, y + h - th)


_REGISTERED: tuple[str, str] | None = None


def _reportlab_font() -> tuple[str, str]:
    """Registra in reportlab lo stesso font di sistema usato per l'HTML."""
    global _REGISTERED
    if _REGISTERED:
        return _REGISTERED
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    _REGISTERED = ("Helvetica", "Helvetica-Bold")
    for candidate in _FONT_CANDIDATES:
        if Path(candidate).is_file():
            bold = candidate.replace("arial.ttf", "arialbd.ttf").replace("Arial.ttf", "Arial Bold.ttf")
            bold = bold.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            pdfmetrics.registerFont(TTFont("Corpo", candidate))
            if Path(bold).is_file():
                pdfmetrics.registerFont(TTFont("Corpo-Bold", bold))
                _REGISTERED = ("Corpo", "Corpo-Bold")
            else:
                _REGISTERED = ("Corpo", "Corpo")
            break
    return _REGISTERED


def eml_to_pdf(raw: bytes, attachments: bool = False) -> Converted:
    """Intestazione, corpo e elenco allegati; con attachments=True accoda gli allegati convertibili."""
    msg = email.message_from_bytes(raw, policy=email.policy.default)
    header_rows = []
    for key, label in (("From", "Da"), ("To", "A"), ("Cc", "Cc"), ("Date", "Data"), ("Subject", "Oggetto")):
        if msg.get(key):
            header_rows.append(f"<tr><td><b>{label}</b></td><td>{_esc(msg.get(key))}</td></tr>")
    body_part = msg.get_body(preferencelist=("html", "plain"))
    inline: dict[str, str] = {}
    files: list[tuple[str, bytes]] = []
    for part in msg.iter_attachments():
        raw_payload = part.get_payload(decode=True)
        payload = raw_payload if isinstance(raw_payload, bytes) else b""
        cid = (part.get("Content-ID") or "").strip("<>")
        if cid and part.get_content_maintype() == "image":
            inline[cid] = f"data:{part.get_content_type()};base64,{base64.b64encode(payload).decode()}"
            continue
        files.append((part.get_filename() or "allegato", payload))
    body = ""
    if body_part is not None:
        content = body_part.get_content()
        if body_part.get_content_subtype() == "html":
            body = re.sub(r"(?is)<(script|style).*?</\1>", "", content)
            body = re.sub(r"cid:([^\"'>\s]+)", lambda m: inline.get(m.group(1), m.group(0)), body)
            body = re.sub(r"(?is)^.*?<body[^>]*>|</body>.*$", "", body)
        else:
            body = text_to_html(content)
    listing = ""
    if files:
        items = "".join(f"<li>{_esc(name)} ({len(data) // 1024 + 1} KB)</li>" for name, data in files)
        listing = f"<h3>Allegati</h3><ul>{items}</ul>"
    html = f"<table class='meta'>{''.join(header_rows)}</table><hr/>{body}{listing}"
    pdf = html_to_pdf(html, None)  # un'email non deve poter leggere file del computer
    result = Converted(pdf)
    if attachments and files:
        result = _append_attachments(result, files)
    return result


def _append_attachments(base: Converted, files: list[tuple[str, bytes]]) -> Converted:
    import tempfile

    from wmt_pdf.ops import Part, merge

    parts = [Part("email", base.data)]
    with tempfile.TemporaryDirectory() as tmp:
        for i, (name, data) in enumerate(files, start=1):
            clean = re.sub(r"[^\w.\- ]", "_", Path(name.replace("\\", "/")).name).strip(". ")
            # Il numero davanti evita doppioni e nomi riservati di Windows (CON, NUL, ...).
            safe = Path(f"{i:02d}_{clean or 'allegato'}")
            target = Path(tmp) / safe
            target.write_bytes(data)
            try:
                kind_of(target)
                # Un allegato non è fidato: niente letture di file locali.
                parts.append(Part(Path(clean).stem or "allegato", to_pdf(target, local_files=False).data))
            except CliError as exc:
                base.notes.append(t("allegato {name} non accodato: {exc}", name=name, exc=exc))
        out = Path(tmp) / "unito.pdf"
        merge(parts, out)
        return Converted(out.read_bytes(), base.notes)
