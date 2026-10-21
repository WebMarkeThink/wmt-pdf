"""Esportazione da PDF verso Word, Excel, HTML, Markdown, testo e immagini."""

from __future__ import annotations

import html as htmllib
import json
import re
from pathlib import Path

from wmt_pdf.extract import Extractor, PageText
from wmt_pdf.i18n import t
from wmt_pdf.source import CliError, Source

TEXT_FORMATS = {"docx", "xlsx", "html", "md", "txt", "json", "xml"}
IMAGE_FORMATS = {"png", "jpg", "tiff"}

# Caratteri che XML (quindi Word ed Excel) non accetta: capitano nel testo estratto dai PDF.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def xml_safe(text: str) -> str:
    return _CONTROL.sub("", text)


_NUMBER = re.compile(r"^-?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?$")


def parse_number(value: str):
    """'1.500,50' → 1500.5; codici con zeri iniziali, percentuali e testo restano stringhe."""
    text = value.strip().replace("€", "").strip()
    if not _NUMBER.match(text) or (len(text) > 1 and text.startswith("0") and not text.startswith("0,")):
        return value
    number = float(text.replace(".", "").replace(",", "."))
    return int(number) if number.is_integer() and "," not in text else number


def extract_pages(source: Source, password: str | None, ocr: str = "auto") -> list[PageText]:
    doc = source.open(password)
    extractor = Extractor(doc, source.data, password=password)
    try:
        pages, _ = extractor.pages(list(range(1, len(doc) + 1)), ocr=ocr)
    finally:
        extractor.close()
        doc.close()
    return pages


def export(source: Source, fmt: str, target: Path, *, password: str | None = None, dpi: int = 150,
           ocr: str = "auto") -> list[Path]:
    if fmt in IMAGE_FORMATS:
        return export_images(source, fmt, target, dpi=dpi, password=password)
    if fmt not in TEXT_FORMATS:
        raise CliError(t("formato di esportazione sconosciuto: {fmt}", fmt=fmt))
    target.parent.mkdir(parents=True, exist_ok=True)
    if fmt in ("json", "xml"):
        data = structure(source, password, ocr)
        if fmt == "json":
            target.write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        else:
            _to_xml(data).write(target, encoding="utf-8", xml_declaration=True)
        return [target]
    pages = extract_pages(source, password, ocr)
    if fmt == "docx":
        _to_docx(pages, target, source.name)
    elif fmt == "xlsx":
        _to_xlsx(pages, target)
    elif fmt == "html":
        target.write_text(_to_html(pages, source.name), encoding="utf-8")
    elif fmt == "md":
        target.write_text(_to_md(pages, source.name), encoding="utf-8")
    else:
        target.write_text("\n\n\f".join(p.text for p in pages) + "\n", encoding="utf-8")
    return [target]


def _segments(doc, number: int) -> list[dict]:
    """Pezzi di testo con il riquadro in punti PDF (origine in basso a sinistra), come li vede PDFium."""
    page = doc[number - 1]
    textpage = page.get_textpage()
    try:
        result = []
        for i in range(textpage.count_rects()):
            left, bottom, right, top = textpage.get_rect(i)
            text = textpage.get_text_bounded(left, bottom, right, top).strip()
            if text:
                result.append({"text": text, "bbox": [round(left, 2), round(bottom, 2), round(right, 2), round(top, 2)]})
        return result
    finally:
        textpage.close()
        page.close()


def structure(source: Source, password: str | None = None, ocr: str = "auto") -> dict:
    """Tutto il documento in una struttura stabile: metadati, pagine (testo, righe, tabelle), campi, annotazioni."""
    from wmt_pdf.forms import list_fields
    from wmt_pdf.marks import list_annotations
    from wmt_pdf.ops import open_pdf, read_outline
    from wmt_pdf.pades import count_signatures

    pages = extract_pages(source, password, ocr)
    doc = source.open(password)
    try:
        meta = doc.get_metadata_dict(skip_empty=True)
        page_data = []
        for item in pages:
            width, height = doc.get_page_size(item.page - 1)
            page_data.append({
                "number": item.page, "width": round(width, 2), "height": round(height, 2),
                "text": item.text, "ocr": item.ocr, "images": item.images, "note": item.note,
                "tables": item.tables, "lines": _segments(doc, item.page) if not item.ocr else [],
            })
    finally:
        doc.close()
    with open_pdf(source.data, password) as pdf:
        outline = read_outline(pdf)

    def tree(items):
        return [{"title": o.title, "page": o.page + 1 if o.page is not None else None, "children": tree(o.children)}
                for o in items]

    return {
        "file": source.name,
        "page_count": len(pages),
        "metadata": meta,
        "pages": page_data,
        "fields": [f.__dict__ for f in list_fields(source.data, password)],
        "annotations": [a.__dict__ for a in list_annotations(source.data, password=password)],
        "outline": tree(outline),
        "signatures": count_signatures(source.data),
    }


def _to_xml(data: dict):
    import xml.etree.ElementTree as ET

    def text(value) -> str:
        return xml_safe("" if value is None else str(value))

    root = ET.Element("document", file=text(data["file"]), page_count=str(data["page_count"]),
                      signatures=str(data["signatures"]))
    meta = ET.SubElement(root, "metadata")
    for name, value in data["metadata"].items():
        ET.SubElement(meta, "item", name=text(name)).text = text(value)
    for page in data["pages"]:
        node = ET.SubElement(root, "page", number=str(page["number"]), width=str(page["width"]),
                             height=str(page["height"]), ocr=str(page["ocr"]).lower(), images=str(page["images"]))
        ET.SubElement(node, "text").text = text(page["text"])
        lines = ET.SubElement(node, "lines")
        for seg in page["lines"]:
            ET.SubElement(lines, "line", bbox=" ".join(str(v) for v in seg["bbox"])).text = text(seg["text"])
        for table in page["tables"]:
            tnode = ET.SubElement(node, "table")
            for row in table:
                rnode = ET.SubElement(tnode, "row")
                for cell in row:
                    ET.SubElement(rnode, "cell").text = text(cell)
    fields = ET.SubElement(root, "fields")
    for f in data["fields"]:
        fnode = ET.SubElement(fields, "field", name=text(f["name"]), type=text(f["type"]), page=text(f["page"]))
        fnode.text = text(f["value"])
    annots = ET.SubElement(root, "annotations")
    for a in data["annotations"]:
        anode = ET.SubElement(annots, "annotation", type=text(a["type"]), page=str(a["page"]))
        anode.text = text(a["contents"] or a["text"])

    def outline(parent, items):
        for item in items:
            onode = ET.SubElement(parent, "bookmark", title=text(item["title"]), page=text(item["page"]))
            outline(onode, item["children"])

    outline(ET.SubElement(root, "outline"), data["outline"])
    ET.indent(root)
    return ET.ElementTree(root)


def _blocks(page: PageText):
    """Il testo della pagina diviso in paragrafi e tabelle Markdown, nell'ordine."""
    buffer: list[str] = []
    table: list[str] = []
    for line in page.text.split("\n"):
        if line.startswith("|"):
            if buffer:
                yield "text", "\n".join(buffer)
                buffer = []
            table.append(line)
            continue
        if table:
            yield "table", table
            table = []
        buffer.append(line)
    if table:
        yield "table", table
    if buffer and "\n".join(buffer).strip():
        yield "text", "\n".join(buffer)


def _md_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if re.match(r"^\|(-+\|)+$", line.replace(" ", "")):
            continue
        cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        rows.append(cells)
    return rows


def _to_docx(pages: list[PageText], target: Path, title: str) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = title
    for i, page in enumerate(pages):
        if i:
            d.add_page_break()
        for kind, block in _blocks(page):
            if kind == "text":
                for para in re.split(r"\n\s*\n", block.strip()):
                    if para.strip():
                        d.add_paragraph(xml_safe(para.replace("\n", " ").strip()))
            else:
                rows = _md_rows(block)
                width = max(len(r) for r in rows)
                table = d.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = xml_safe(value)
        if page.note:
            d.add_paragraph(xml_safe(f"[{page.note}]"))
    d.save(target)


def _to_xlsx(pages: list[PageText], target: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for page in pages:
        for k, rows in enumerate(page.tables, start=1):
            ws = wb.create_sheet(t("Pagina {page}", page=page.page) + (f" ({k})" if len(page.tables) > 1 else ""))
            for row in rows:
                ws.append([parse_number(xml_safe(c)) for c in row])
    if not wb.worksheets:
        raise CliError(t("nessuna tabella nel PDF da esportare in Excel: prova --to docx o md"))
    wb.save(target)


def _to_html(pages: list[PageText], title: str) -> str:
    out = [f"<!doctype html><html lang='it'><head><meta charset='utf-8'><title>{htmllib.escape(title)}</title>",
           ("<style>body{font-family:sans-serif;max-width:50em;margin:2em auto;line-height:1.5}"
           "table{border-collapse:collapse}td,th{border:1px solid #999;padding:4px 8px}</style></head><body>")]
    for page in pages:
        out.append(t("<section><h2>Pagina {page}</h2>", page=page.page))
        for kind, block in _blocks(page):
            if kind == "text":
                for para in re.split(r"\n\s*\n", block.strip()):
                    if para.strip():
                        out.append("<p>" + "<br>".join(htmllib.escape(x) for x in para.split("\n")) + "</p>")
            else:
                rows = _md_rows(block)
                out.append("<table>")
                for r, row in enumerate(rows):
                    tag = "th" if r == 0 else "td"
                    out.append("<tr>" + "".join(f"<{tag}>{htmllib.escape(c)}</{tag}>" for c in row) + "</tr>")
                out.append("</table>")
        out.append("</section>")
    out.append("</body></html>")
    return "\n".join(out)


def _to_md(pages: list[PageText], title: str) -> str:
    parts = [f"# {title}"]
    for page in pages:
        parts.append(t("\n## Pagina {page}\n\n{text}", page=page.page, text=page.text))
    return "\n".join(parts) + "\n"


def single_file(fmt: str, target: Path) -> bool:
    """tiff verso un file .tif/.tiff è un unico file multipagina; png e jpg vanno in una cartella."""
    return fmt not in IMAGE_FORMATS or (fmt == "tiff" and target.suffix.lower() in (".tif", ".tiff"))


def image_paths(source: Source, fmt: str, target: Path, count: int) -> list[Path]:
    stem = Path(source.name).stem.removesuffix(".pdf")
    return [target / f"{stem}_p{n}.{fmt}" for n in range(1, count + 1)]


def _render(doc, n: int, dpi: int):
    page = doc[n - 1]
    try:
        return page.render(scale=dpi / 72).to_pil().convert("RGB")
    finally:
        page.close()


def export_images(source: Source, fmt: str, target: Path, *, dpi: int, password: str | None) -> list[Path]:
    doc = source.open(password)
    doc.init_forms()  # senza, PDFium non disegna i campi dei moduli
    try:
        count = len(doc)
        if single_file(fmt, target):
            target.parent.mkdir(parents=True, exist_ok=True)
            # Il TIFF multipagina vuole tutte le immagini insieme: si generano una alla volta.
            first = _render(doc, 1, dpi)
            rest = (_render(doc, n, dpi) for n in range(2, count + 1))
            first.save(target, save_all=True, append_images=rest, compression="tiff_deflate", dpi=(dpi, dpi))
            return [target]
        target.mkdir(parents=True, exist_ok=True)
        written = []
        for n, path in enumerate(image_paths(source, fmt, target, count), start=1):
            options = {"quality": 85} if fmt == "jpg" else {}
            _render(doc, n, dpi).save(path, dpi=(dpi, dpi), **options)  # una pagina alla volta in memoria
            written.append(path)
        return written
    finally:
        doc.close()
